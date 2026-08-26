"""
Skrip admin: impor data tabungan dari Excel rekap ke Google Sheet Tabungan.

    python3 tab_import.py smp "TABUNGAN SMPIA data mas oki.xlsx"
    python3 tab_import.py smp "file.xlsx" --dry-run     # cek saja, tanpa menulis

BENTUK FILE SUMBER (per sheet = satu kelas, nama sheet "7"/"8"/"9"):
  Baris 2 : label tahun ajaran di atas tiap blok ("2024/2025", ...) + kolom
            ringkasan (JUMLAH PENYETORAN / PENGAMBILAN / SALDO)
  Baris 3 : nama bulan (JULI ..) dan "TOTAL" di kolom ke-13 tiap blok
  Baris 4+: data siswa — kolom A URUT, B INDUK, C NAMA, D KELAS
  Tiap blok tahun = 13 kolom: 12 bulan + 1 TOTAL.

Nilai bulanan = SETORAN bulan itu (bukan saldo berjalan). PENGAMBILAN hanya
dicatat per TAHUN, tanpa bulan/tanggal — di sheet ditaruh di bulan JUNI
(bulan terakhir tahun ajaran). Tanggal setor/tarik dibiarkan kosong karena
file sumber memang tidak menyimpannya; kolom SALDO tidak butuh tanggal.

Yang dijadikan acuan adalah KOLOM BULANAN, bukan kolom TOTAL/SALDO — di file
contoh ada baris yang sel TOTAL-nya tertimpa angka 0 sehingga setoran
bulanannya tidak terhitung. Mengimpor dari kolom bulanan sekaligus
membetulkan itu.

SALDO AWAL: tahun paling awal = 0, tahun berikutnya memakai formula rantai ke
tab tahun sebelumnya (lihat tab_sheet.link_saldo_awal).
"""
import sys

import openpyxl

import tab_config as C
import tab_sheet as SL

MONTHS = 12
BLOCK_COLS = MONTHS + 1        # 12 bulan + kolom TOTAL
COL_INDUK, COL_NAMA = 2, 3
FIRST_BLOCK_COL = 5


def _num(v):
    if v in (None, ""):
        return 0
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(".", "").replace(",", "."))
    except ValueError:
        return 0


def read_source(path, sheet):
    """-> (list blok tahun, list siswa). Tiap siswa: induk, nama, per-tahun
    daftar setoran bulanan + pengambilan tahunan."""
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]

    blocks, col = [], FIRST_BLOCK_COL
    while True:
        label = ws.cell(2, col).value
        if not (label and "/" in str(label)):
            break
        blocks.append({"label": str(label).strip(), "start": col,
                       "year": int(str(label).split("/")[0])})
        col += BLOCK_COLS

    ambil_col = next((c for c in range(1, ws.max_column + 1)
                      if ws.cell(2, c).value == "PENGAMBILAN"), None)

    siswa = []
    for r in range(4, ws.max_row + 1):
        induk = ws.cell(r, COL_INDUK).value
        if induk in (None, ""):
            continue
        rec = {"induk": str(induk).strip(),
               "nama": str(ws.cell(r, COL_NAMA).value or "").strip(), "years": {}}
        for i, b in enumerate(blocks):
            rec["years"][b["year"]] = {
                "setor": [_num(ws.cell(r, b["start"] + m).value) for m in range(MONTHS)],
                "tarik": _num(ws.cell(r, ambil_col + i).value) if ambil_col else 0,
            }
        siswa.append(rec)
    return blocks, siswa


def pair_rows(roster, siswa, title):
    """Pasangkan siswa di file dengan baris di sheet BERDASARKAN URUTAN.

    Sengaja tidak memakai INDUK sebagai kunci: di data kelas 9 ada satu INDUK
    yang dipakai dua siswa berbeda (0229 — RADITYA & RAISA), sehingga pemetaan
    lewat INDUK membuat salah satunya menimpa yang lain. Urutan dipakai sebagai
    kunci, tapi INDUK tiap posisi tetap dicocokkan supaya pergeseran urutan
    langsung ketahuan, bukan diam-diam salah tulis."""
    if len(roster) != len(siswa):
        raise SystemExit(f"{title}: jumlah siswa beda (sheet {len(roster)}, "
                         f"file {len(siswa)}) — impor dibatalkan.")
    salah = [(i, r["induk"], s["induk"])
             for i, (r, s) in enumerate(zip(roster, siswa)) if r["induk"] != s["induk"]]
    if salah:
        raise SystemExit(f"{title}: urutan siswa tidak sama dengan file "
                         f"(mis. posisi {salah[0][0]}: sheet {salah[0][1]} vs "
                         f"file {salah[0][2]}) — impor dibatalkan.")
    return [(r["row"], s) for r, s in zip(roster, siswa)]


def build_grid(pairs, year):
    """Susun nilai kolom E..(kolom terakhir) untuk seluruh baris data tab."""
    months = C.months_for_year(2024)          # selalu 12 bulan Juli..Juni
    n_cols = len(months) * C.COLS_PER_MONTH
    rows = [p[0] for p in pairs]
    first, last = rows[0], rows[-1]
    by_row = dict(pairs)

    grid = []
    for row in range(first, last + 1):
        cells = ["" for _ in range(n_cols)]
        rec = by_row.get(row)
        for pos in range(len(months)):
            base = pos * C.COLS_PER_MONTH
            if rec:
                y = rec["years"].get(year)
                if y:
                    if y["setor"][pos]:
                        cells[base + 1] = y["setor"][pos]        # SETOR
                    if pos == MONTHS - 1 and y["tarik"]:
                        cells[base + 3] = y["tarik"]             # TARIK di Juni
            # SALDO selalu formula berjalan
            cells[base + 4] = C.saldo_formula(row, pos)
        grid.append(cells)
    return grid, first, last


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    if len(args) < 2:
        print(__doc__)
        sys.exit(1)
    pkey, path = args[0].lower(), args[1]
    if pkey not in C.PROFILES:
        print("Jenjang harus salah satu dari:", ", ".join(C.PROFILES))
        sys.exit(1)

    prof = C.PROFILES[pkey]
    jenjang, kelas_list = prof["jenjang"], prof["kelas"]
    book = SL.open_book(prof["spreadsheet_id"])
    print("Terhubung:", book.title, "| dry-run" if dry else "")

    ringkas = []
    for kelas in kelas_list:
        try:
            blocks, siswa = read_source(path, str(kelas))
        except KeyError:
            print(f"  ! sheet '{kelas}' tidak ada di file sumber, dilewati")
            continue
        print(f"\nKELAS {kelas}: {len(siswa)} siswa, blok "
              f"{[b['label'] for b in blocks]}")

        for b in blocks:
            year = b["year"]
            title = C.tab_name(kelas, year, jenjang)
            try:
                ws = book.worksheet(title)
            except Exception:
                print(f"  ! tab {title} belum ada — jalankan tab_backfill.py dulu")
                continue
            roster = SL.read_roster_from_tab(ws)
            pairs = pair_rows(roster, siswa, title)
            grid, first, last = build_grid(pairs, year)

            setor = sum(v for rec in siswa for v in rec["years"].get(year, {}).get("setor", []))
            tarik = sum(rec["years"].get(year, {}).get("tarik", 0) for rec in siswa)
            ringkas.append((title, setor, tarik))
            print(f"  {title}: setor Rp {_rp(setor)} | tarik Rp {_rp(tarik)}")

            if dry:
                continue
            start_a1 = SL.rowcol_to_a1(first, C.FIRST_MONTH_COL)
            end_a1 = SL.rowcol_to_a1(last, C.FIRST_MONTH_COL + len(grid[0]) - 1)
            ws.update(grid, f"{start_a1}:{end_a1}", value_input_option="USER_ENTERED")

        # SALDO AWAL: tahun paling awal 0, sisanya dirantai ke tahun sebelumnya
        if dry:
            continue
        years_tab = SL.list_years(book, jenjang)
        for i, year in enumerate(years_tab):
            title = C.tab_name(kelas, year, jenjang)
            try:
                ws = book.worksheet(title)
            except Exception:
                continue
            if i == 0:
                roster = SL.read_roster_from_tab(ws)
                if roster:
                    f, l = roster[0]["row"], roster[-1]["row"]
                    ws.update([[0] for _ in range(f, l + 1)], f"D{f}:D{l}",
                              value_input_option="USER_ENTERED")
            else:
                prev_title = C.tab_name(kelas, years_tab[i - 1], jenjang)
                SL.link_saldo_awal(ws, book.worksheet(prev_title), prev_title,
                                   years_tab[i - 1])
        print(f"  SALDO AWAL kelas {kelas}: {C.tab_name(kelas, years_tab[0], jenjang)} = 0, "
              f"tahun berikutnya dirantai")

    print("\nRINGKASAN")
    for title, setor, tarik in ringkas:
        print(f"  {title:<14} setor Rp {_rp(setor):>14}  tarik Rp {_rp(tarik):>14}")
    if not dry:
        print("\nSelesai:", f"https://docs.google.com/spreadsheets/d/{prof['spreadsheet_id']}")


def _rp(n):
    return format(int(round(n)), ",").replace(",", ".")


if __name__ == "__main__":
    main()
