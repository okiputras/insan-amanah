"""
Skrip admin: impor master siswa & data pemasukan BPP dari Excel sekolah ke spreadsheet
"PEMBAYARAN SD INSAN AMANAH" (menu Pembayaran SD).

    python3 bayar_import.py data-alarm --dry-run      # cek dulu, tanpa menulis
    python3 bayar_import.py data-alarm                # tulis ke Google Sheet
    python3 bayar_import.py data-alarm --isi-ulang-siswa   # timpa tab SISWA dari master
    python3 bayar_import.py data-alarm --perbaiki     # samakan nominal baris lama dengan sumber

Folder berisi:
  TEMPLATE NOMINAL BPP DAN KEGIATAN SDIA <T.A.>.xlsx  -> master (sheet KELAS 1..6, tiap
      sheet berisi beberapa blok rombel bertumpuk: baris label "1A", header NO|INDUK|NAMA
      SISWA|BPP|KATERING|KEGIATAN|TOTAL, data, baris TOTAL)
  PENGELUARAN PEMASUKAN SDIA <BULAN> <TAHUN> ...xlsx  -> hanya tab "PEMASUKAN BPP" dipakai

TAB PEMASUKAN BPP tersusun horizontal: satu BLOK per tanggal terima, berdampingan.
Baris 3: label tanggal di kolom pertama blok, "NAMA" di kolom kedua, nama bulan di
kolom ketiga. Baris 4: sub-header. Dua format blok:
  8 kolom: INDUK | NAMA | BPP | KATERING | KEG | TABUNGAN | TOTAL
  9 kolom: INDUK | NAMA | BPP+KATERING | BPP | KATERING | KEG | TABUNGAN | TOTAL
           (kolom "BPP" pertama adalah gabungan; judul "BPP" kedua kadang kosong, jadi
           format dibedakan dari posisi judul TOTAL — lihat _lebar9)
Blok kosong (sisa templat) dilewati. INDUK kadang angka, kadang teks, dan di file Juli
sebagian besar kosong — di situ siswa dicocokkan lewat nama.

Baris yang tidak bisa dipastikan siswanya TIDAK ditebak: masuk tab PERLU DICEK bersama
saran, untuk diputuskan manusia di aplikasi. Setiap baris membawa KUNCI
"<BULAN TAHUN>:<kolom blok>:<baris>", jadi skrip aman dijalankan ulang tanpa
menggandakan uang.
"""
import datetime as dt
import difflib
import glob
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

import bayar_config as C

PINDAH = "PINDAH"


# ---------------------------------------------------------------- master
def baca_master(path):
    """-> list siswa (urut master). Koreksi yang dikonfirmasi sekolah diterapkan di sini
    dan dicatat di kolom CATATAN, supaya tab SISWA jadi sumber yang benar."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for name in wb.sheetnames:
        ws = wb[name]
        rombel = None
        for r in range(1, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if isinstance(a, str) and re.fullmatch(r"\s*\d[A-Z]\s*", a) and b is None:
                rombel = a.strip()
                continue
            if not (isinstance(a, (int, float)) and b is not None and rombel):
                continue
            v = [ws.cell(r, c).value for c in range(1, 8)]
            pindah = any(isinstance(x, str) and PINDAH in x.upper() for x in v[3:6])
            num = lambda x: int(x) if isinstance(x, (int, float)) else 0
            s = {"induk": str(int(b)), "nama": str(v[2]).strip(), "rombel": rombel,
                 "bpp": num(v[3]), "katering": num(v[4]), "kegiatan": num(v[5]),
                 "status": C.STATUS_PINDAH if pindah else C.STATUS_AKTIF, "catatan": ""}
            if pindah:
                s["catatan"] = "master: bertanda PINDAH"
            if rombel.startswith("4") and s["katering"] == 220_000:
                s["katering"] = 200_000
                s["catatan"] = "katering dikoreksi dari 220.000 (master salah, dikonfirmasi sekolah)"
            out.append(s)
    return out


# ---------------------------------------------------------------- pembayaran
def _induk(v):
    if isinstance(v, (int, float)):
        return str(int(v))
    if isinstance(v, str) and re.fullmatch(r"\s*\d{3,5}\s*", v):
        return v.strip()
    return None


def _uang(v):
    return int(round(v)) if isinstance(v, (int, float)) else 0


def _bulan_file(path):
    m = re.search(r"SDIA\s+([A-Z]+)\s+(\d{4})", os.path.basename(path).upper())
    if not m or m.group(1) not in C.MONTHS_ID:
        raise ValueError(f"Tidak bisa membaca bulan dari nama file: {path}")
    return C.MONTHS_ID.index(m.group(1)) + 1, int(m.group(2)), f"{m.group(1)} {m.group(2)}"


def _tanggal(label, bulan, tahun):
    """Label blok -> tanggal ISO. "1-17 JULI" -> tanggal akhir rentang; tanggal Excel ->
    apa adanya; "TRANSFER" (tanpa tanggal) -> akhir bulan file."""
    if isinstance(label, dt.datetime):
        return label.date().isoformat()
    s = str(label or "").upper()
    m = re.search(r"(\d{1,2})\s*(?:-\s*(\d{1,2}))?\s+([A-Z]+)(?:\s+(\d{4}))?", s)
    if m and m.group(3) in C.MONTHS_ID:
        hari = int(m.group(2) or m.group(1))
        bln = C.MONTHS_ID.index(m.group(3)) + 1
        th = int(m.group(4) or tahun)
        try:
            return dt.date(th, bln, hari).isoformat()
        except ValueError:
            pass
    akhir = (dt.date(tahun + (bulan == 12), bulan % 12 + 1, 1) - dt.timedelta(days=1))
    return akhir.isoformat()


def _lebar9(ws, c):
    """Blok format 9 kolom? Ditentukan dari POSISI judul TOTAL di baris 4 (kolom ke-8 blok
    = 9 kolom, ke-7 = 8 kolom). Jangan dari teks "BPP" ganda: di blok 8 Agustus judul BPP
    kedua kosong, sehingga blok itu terbaca 8 kolom dan semua nominal bergeser satu kolom
    (BPP+katering terhitung dua kali, kegiatan masuk ke tabungan)."""
    for k in (7, 6):
        if str(ws.cell(4, c + k).value or "").strip().upper().startswith("TOTAL"):
            return k == 7
    return [str(ws.cell(4, c + k).value or "").strip().upper() for k in (2, 3)] == ["BPP", "BPP"]


def baca_pembayaran(path):
    """-> (baris pembayaran, info blok untuk rekonsiliasi)."""
    bulan, tahun, label_file = _bulan_file(path)
    ws = openpyxl.load_workbook(path, data_only=True)["PEMASUKAN BPP"]
    blok = [c for c in range(1, ws.max_column + 1)
            if str(ws.cell(3, c + 1).value or "").strip().upper() == "NAMA"]
    rows, info = [], []
    for c in blok:
        lebar9 = _lebar9(ws, c)
        label = ws.cell(3, c).value
        tanggal = _tanggal(label, bulan, tahun)
        label_txt = label.date().isoformat() if isinstance(label, dt.datetime) else str(label or "").strip()
        terakhir = None
        jumlah_bpp = 0
        janggal = []
        for r in range(5, ws.max_row + 1):
            nama = ws.cell(r, c + 1).value
            if nama in (None, ""):
                continue
            v = [ws.cell(r, c + k).value for k in range(8)]
            if lebar9:
                gabung, kat, keg, tab = _uang(v[2]), _uang(v[4]), _uang(v[5]), _uang(v[6])
                # BPP diturunkan dari kolom gabungan (diketik dari nominal bayar), bukan dibaca
                # dari kolom "BPP": kolom itu rumus yang kadang memakai konstanta salah
                # (=900000-katering untuk siswa bertagihan 855.000). TOTAL juga tidak dipakai:
                # di blok 6 September rumus TOTAL sekolah ikut menjumlah kolom gabungan.
                bpp = gabung - kat if gabung else _uang(v[3])
                if _uang(v[3]) != bpp:
                    janggal.append(r)
            else:
                bpp, kat, keg, tab = _uang(v[2]), _uang(v[3]), _uang(v[4]), _uang(v[5])
                if _uang(v[6]) != bpp + kat + keg + tab:
                    janggal.append(r)
            terakhir = r
            jumlah_bpp += bpp
            rows.append({"tanggal": tanggal, "induk": _induk(v[0]), "nama": str(nama).strip(),
                         "bpp": bpp, "katering": kat, "kegiatan": keg, "tabungan": tab,
                         "kunci": f"{label_file}:{get_column_letter(c)}:{r}",
                         "label": label_txt})
        if terakhir is None:
            continue
        kol_bpp = c + 3 if lebar9 else c + 2
        footer = next((ws.cell(r, kol_bpp).value for r in range(terakhir + 1, terakhir + 4)
                       if isinstance(ws.cell(r, kol_bpp).value, (int, float))
                       and ws.cell(r, c + 1).value in (None, "")), None)
        info.append({"blok": get_column_letter(c), "label": label_txt, "lebar9": lebar9,
                     "jumlah": jumlah_bpp, "footer": footer, "janggal": janggal})
    return label_file, rows, info


# ---------------------------------------------------------------- cocokkan nama
def norm(s):
    s = re.sub(r"\(.*?\)", " ", str(s).upper())     # buang keterangan "(JUL - DES)"
    s = re.sub(r"[^A-Z ]", " ", s)                   # titik, apostrof, tanda hubung
    return re.sub(r"\s+", " ", s).strip()


def _kata_cocok(t, w):
    """Satu kata dari nama-pembayaran cocok dengan satu kata nama-master?"""
    if t == w:
        return True
    if len(t) == 1:                                   # inisial: A = AHMAD, M = MUHAMMAD
        return w.startswith(t)
    if len(t) >= 3 and w.startswith(t):               # dipotong: PRAT ~ PRATAMA
        return True
    return len(t) >= 4 and difflib.SequenceMatcher(None, t, w).ratio() >= 0.8   # salah ketik: SAPUTRAH ~ SAPUTRA


def _berurutan(kata, kata_master):
    i = 0
    for t in kata:
        while i < len(kata_master) and not _kata_cocok(t, kata_master[i]):
            i += 1
        if i == len(kata_master):
            return False
        i += 1
    return True


def cocokkan_nama(nama, siswa_norm, siswa, bpp_bayar=0):
    """-> (induk | None, cara, kandidat). Konservatif: bila lebih dari satu siswa cocok
    dan nominal BPP tidak bisa memisahkannya, TIDAK menebak.

    Teruji terhadap 1.239 baris Agustus/September yang INDUK-nya diketahui: 99,5% tepat."""
    n = norm(nama)
    if not n:
        return None, "kosong", []
    kata = n.split()
    hit = [k for k, f in siswa_norm.items() if f == n]
    if len(hit) == 1:
        return hit[0], "persis", hit
    hit = [k for k, f in siswa_norm.items() if _berurutan(kata, f.split())]
    if len(hit) == 1:
        return hit[0], "kata", hit
    if len(hit) > 1 and bpp_bayar:
        sempit = [k for k in hit if siswa[k]["bpp"] and bpp_bayar % siswa[k]["bpp"] == 0]
        if len(sempit) == 1:
            return sempit[0], "kata+nominal", hit
    if len(hit) > 1:
        return None, "ambigu", hit
    skor = sorted(((difflib.SequenceMatcher(None, n, f).ratio(), k) for k, f in siswa_norm.items()),
                  reverse=True)
    if skor and skor[0][0] >= 0.85 and (len(skor) < 2 or skor[0][0] - skor[1][0] >= 0.08):
        return skor[0][1], "mirip", [skor[0][1]]
    return None, "tidak ketemu", [k for _, k in skor[:3]]


_TA_LALU = re.compile(r"\((JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI)\s+%d\)" % C.TAHUN_AJARAN)


def tentukan_siswa(p, siswa, siswa_norm):
    """-> (induk, masalah | None, saran | None, cara).

    INDUK tertulis dipercaya, KECUALI namanya jelas orang lain dan namanya cocok ke siswa
    lain — di data nyata ini menangkap salah ketik (2037 untuk 2307, 2315 untuk 2316)."""
    induk = p["induk"]
    if induk and induk in siswa:
        a, b = norm(p["nama"]), norm(siswa[induk]["nama"])
        mirip = difflib.SequenceMatcher(None, a, b).ratio()
        if mirip < 0.6 and not (a and b.startswith(a.split()[0])):
            lain, _, _ = cocokkan_nama(p["nama"], siswa_norm, siswa, p["bpp"])
            if lain and lain != induk:
                return None, (f"INDUK {induk} milik {siswa[induk]['nama']}, tetapi nama "
                              f"tercatat cocok dengan {lain} {siswa[lain]['nama']}"), lain, "induk"
        if siswa[induk]["status"] == C.STATUS_PINDAH:
            return None, f"{siswa[induk]['nama']} berstatus PINDAH di master tapi tercatat membayar", induk, "induk"
        masalah, cara = None, "induk"
    elif induk:
        k, _, _ = cocokkan_nama(p["nama"], siswa_norm, siswa, p["bpp"])
        return None, f"INDUK {induk} tidak ada di master", k, "induk"
    else:
        k, cara, kand = cocokkan_nama(p["nama"], siswa_norm, siswa, p["bpp"])
        if not k:
            # Sengaja TANPA saran: pencocok menolak menebak di sini, dan saran mengisi
            # kotak input di aplikasi — satu klik Simpan tanpa membaca akan mengkreditkan
            # uang ke siswa yang salah (pernah terjadi: nama dua kata yang tercatat
            # terpotong, kandidat pertamanya siswa lain yang ejaannya mirip, padahal
            # siswa yang benar bernama tiga kata). Kandidat tetap disebut di MASALAH.
            daftar = ", ".join(f"{x} {siswa[x]['nama']}" for x in kand[:3])
            return None, f"nama {cara} — kandidat: {daftar}", None, "nama"
        induk, masalah = k, None
    if _TA_LALU.search(p["nama"].upper()):
        return None, "kemungkinan pembayaran tunggakan tahun ajaran lalu", induk, cara
    return induk, masalah, None, cara


# ---------------------------------------------------------------- main
def _rp(n):
    return format(int(n), ",").replace(",", ".")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    isi_ulang = "--isi-ulang-siswa" in sys.argv
    perbaiki = "--perbaiki" in sys.argv
    if not args:
        print(__doc__)
        sys.exit(1)
    folder = args[0]
    master_path = next(iter(sorted(glob.glob(os.path.join(folder, "TEMPLATE NOMINAL*.xlsx")))), None)
    bayar_paths = sorted(glob.glob(os.path.join(folder, "PENGELUARAN PEMASUKAN SDIA*.xlsx")),
                         key=lambda p: (_bulan_file(p)[1], C.bulan_list().index(_bulan_file(p)[0])))
    if not master_path or not bayar_paths:
        print("Folder harus berisi file TEMPLATE NOMINAL...xlsx dan PENGELUARAN PEMASUKAN SDIA...xlsx")
        sys.exit(1)

    daftar_siswa = baca_master(master_path)
    siswa = {s["induk"]: s for s in daftar_siswa}
    siswa_norm = {k: norm(s["nama"]) for k, s in siswa.items()}
    n_pindah = sum(s["status"] == C.STATUS_PINDAH for s in daftar_siswa)
    n_kor = sum("dikoreksi" in s["catatan"] for s in daftar_siswa)
    print(f"Master: {len(daftar_siswa)} siswa, {n_pindah} PINDAH, {n_kor} katering dikoreksi 220.000 -> 200.000")

    book = None
    terpakai = {}
    if not dry and not C.SPREADSHEET_ID:
        print("SPREADSHEET_ID belum diisi di bayar_config.py.")
        sys.exit(1)
    if C.SPREADSHEET_ID:
        import gspread
        import bayar_sheet as S
        book = S.open_book()
        if not dry:                        # dry-run tidak boleh menulis apa pun, termasuk membuat tab
            S.siapkan_tab(book)
        try:
            terpakai = S.nominal_per_kunci(book)
        except gspread.exceptions.WorksheetNotFound:
            terpakai = {}                  # dry-run sebelum impor pertama: tab belum ada
        print("Terhubung:", book.title, "| baris yang sudah ada:", len(terpakai))

    masuk, review, blok_beda, blok_janggal, beda = [], [], [], [], []
    for path in bayar_paths:
        bln, th, label_file = _bulan_file(path)
        if (th, bln) >= (C.LAPORAN_VA_MULAI.year, C.LAPORAN_VA_MULAI.month):
            print(f"  {label_file:<15} DILEWATI — sejak {C.label_bulan(C.LAPORAN_VA_MULAI.month)} pembayaran VA "
                  "dicatat dari laporan bank (Data Validasi SD) dan transfer diinput di aplikasi; "
                  "mengimpor Excel bulan ini akan mendobelkan uang.")
            continue
        label_file, rows, info = baca_pembayaran(path)
        n_induk = n_nama = n_rev = n_lewat = 0
        for p in rows:
            if p["kunci"] in terpakai:
                n_lewat += 1
                tab, baris, lama = terpakai[p["kunci"]]
                baru = [p["bpp"], p["katering"], p["kegiatan"], p["tabungan"]]
                if lama != baru:
                    beda.append((tab, baris, p, lama, baru))
                continue
            induk, masalah, saran, cara = tentukan_siswa(p, siswa, siswa_norm)
            if masalah:
                n_rev += 1
                review.append({"tanggal": p["tanggal"], "nama": p["nama"], "induk_tercatat": p["induk"] or "",
                               "bpp": p["bpp"], "katering": p["katering"], "kegiatan": p["kegiatan"],
                               "tabungan": p["tabungan"], "masalah": masalah, "saran": saran or "",
                               "kunci": p["kunci"]})
                continue
            n_induk += cara == "induk"
            n_nama += cara != "induk"
            catatan = f"impor {label_file}, blok {p['label']}"
            if cara != "induk":
                catatan += f"; dicocokkan dari nama '{p['nama']}'"
            masuk.append({"tanggal": p["tanggal"], "induk": induk, "nama": siswa[induk]["nama"],
                          "bpp": p["bpp"], "katering": p["katering"], "kegiatan": p["kegiatan"],
                          "tabungan": p["tabungan"], "sumber": f"impor {label_file}",
                          "kunci": p["kunci"], "catatan": catatan})
        for b in info:
            if b["footer"] is not None and abs(b["footer"] - b["jumlah"]) >= 1:
                blok_beda.append((label_file, b))
            if b["janggal"]:
                blok_janggal.append((label_file, b))
        print(f"  {label_file:<15} {len(rows):>4} baris | via INDUK {n_induk:>4} | via nama {n_nama:>4} "
              f"| perlu dicek {n_rev:>2}" + (f" | sudah ada {n_lewat}" if n_lewat else ""))

    tot = lambda d, k: sum(x[k] for x in d)
    print(f"\nAkan dicatat: {len(masuk)} baris — BPP {_rp(tot(masuk, 'bpp'))}, katering "
          f"{_rp(tot(masuk, 'katering'))}, kegiatan {_rp(tot(masuk, 'kegiatan'))}")
    print(f"Perlu dicek : {len(review)} baris — BPP {_rp(tot(review, 'bpp'))}")
    for r in review:
        print(f"   - {r['nama'][:30]:<30} {r['induk_tercatat'] or '-':<5} BPP {_rp(r['bpp']):>10} | {r['masalah']}")

    if blok_beda:
        print("\nRekonsiliasi: blok yang jumlah barisnya TIDAK sama dengan footer di Excel sekolah")
        print("(ini temuan di file sumber, bukan kegagalan impor — yang diimpor adalah per baris):")
        for label_file, b in blok_beda:
            print(f"   - {label_file} blok {b['blok']} ({b['label']}): jumlah baris BPP "
                  f"{_rp(b['jumlah'])} vs footer {_rp(b['footer'])} (selisih {_rp(b['jumlah'] - b['footer'])})")

    if blok_janggal:
        print("\nBaris yang tidak konsisten di dalam bloknya sendiri (cek manual di Excel sekolah):")
        for label_file, b in blok_janggal:
            print(f"   - {label_file} blok {b['blok']} ({b['label']}, {'9' if b['lebar9'] else '8'} kolom): "
                  f"baris {', '.join(map(str, b['janggal'][:10]))}")

    if beda:
        print(f"\nNominal di Sheet BERBEDA dari file sumber: {len(beda)} baris "
              f"(BPP / katering / kegiatan / tabungan, Sheet -> sumber)")
        for tab, baris, p, lama, baru in beda:
            print(f"   - {p['kunci']:<22} {tab} baris {baris:<5} {p['nama'][:24]:<24} "
                  f"{'/'.join(_rp(x) for x in lama)} -> {'/'.join(_rp(x) for x in baru)}")
        if not perbaiki:
            print("   Jalankan dengan --perbaiki untuk menyamakan dengan sumber (baris yang sengaja"
                  " dikoreksi manual di Sheet ikut tertimpa — periksa daftar di atas dulu).")

    if dry:
        print("\n(dry-run: tidak ada yang ditulis)")
        return
    import bayar_sheet as S
    if isi_ulang or not S.baca_siswa(book):
        S.isi_siswa(book, daftar_siswa)
        print(f"\nTab {C.TAB_SISWA} diisi: {len(daftar_siswa)} siswa")
    S.catat_banyak(book, masuk)
    S.tambah_review(book, review)
    print(f"Ditulis: {len(masuk)} pembayaran, {len(review)} baris Perlu Dicek")
    if perbaiki and beda:
        for tab in (C.TAB_BAYAR, C.TAB_REVIEW):
            S.koreksi_nominal(book, tab, [(baris, baru) for t, baris, _, _, baru in beda if t == tab])
        print(f"Dikoreksi: {len(beda)} baris disamakan dengan sumber")
    print("Selesai:", f"https://docs.google.com/spreadsheets/d/{C.SPREADSHEET_ID}")


if __name__ == "__main__":
    main()
